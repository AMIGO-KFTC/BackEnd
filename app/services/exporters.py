"""인수인계서 Markdown → Word(.docx) / PDF 변환.

Markdown 은 mistune 으로 구문 트리를 만든 뒤 제목·문단·목록·표만 옮긴다(에이전트가 만드는 문서 형식).
셀 병합(rowspan/colspan)이나 막대 색이 필요한 표(월별 일정표, 협업관계 칸 분할)는 문서 안에 HTML <table> 로 들어 있으며,
이 HTML 표도 읽어서 같은 모양의 Word/PDF 표로 옮긴다(th 는 머리글, class="bar"·bgcolor 칸은 색칠).
- Word: 한글이 깨지지 않도록 동아시아 글꼴(eastAsia)을 '맑은 고딕'으로 지정
- PDF : reportlab. 한글 TTF 글꼴을 찾아 문서에 포함(나눔고딕·맑은 고딕·애플고딕·WenQuanYi …).
        못 찾으면 reportlab 내장 CID 글꼴(HYGothic-Medium, 뷰어 글꼴 사용)로 대체. AMIGO_PDF_FONT 로 지정 가능.
"""

from __future__ import annotations

import io
import os
from html.parser import HTMLParser
from dataclasses import dataclass, field
from datetime import date
from xml.sax.saxutils import escape

import mistune

_MD = mistune.create_markdown(renderer=None, plugins=["table", "strikethrough"])


# --------------------------------------------------------------------------- Markdown → 블록
@dataclass
class Run:
    text: str
    bold: bool = False
    italic: bool = False
    code: bool = False


@dataclass
class Block:
    kind: str  # heading | paragraph | list_item | table | rule
    runs: list[Run] = field(default_factory=list)
    level: int = 0  # heading level / list depth
    ordered: bool = False
    number: int = 0
    rows: list[list[list[Run]]] = field(default_factory=list)  # table: rows → cells → runs (Markdown 표는 첫 행이 머리글)
    # HTML 표 전용: (행, 열) → (rowspan, colspan) 은 병합의 왼쪽 위 칸, 머리글 칸 목록(None 이면 첫 행), 칸 색(RRGGBB)
    spans: dict[tuple[int, int], tuple[int, int]] = field(default_factory=dict)
    headers: set[tuple[int, int]] | None = None
    fills: dict[tuple[int, int], str] = field(default_factory=dict)

    def is_header(self, r: int, c: int) -> bool:
        return r == 0 if self.headers is None else (r, c) in self.headers


def _runs(children: list[dict], bold: bool = False, italic: bool = False) -> list[Run]:
    out: list[Run] = []
    for node in children or []:
        kind = node.get("type")
        if kind == "text":
            out.append(Run(node.get("raw", ""), bold, italic))
        elif kind == "strong":
            out.extend(_runs(node.get("children", []), True, italic))
        elif kind == "emphasis":
            out.extend(_runs(node.get("children", []), bold, True))
        elif kind == "codespan":
            out.append(Run(node.get("raw", ""), bold, italic, code=True))
        elif kind in ("softbreak", "linebreak"):
            out.append(Run("\n", bold, italic))
        elif kind == "link":
            out.extend(_runs(node.get("children", []), bold, italic))
        elif "children" in node:
            out.extend(_runs(node["children"], bold, italic))
        elif "raw" in node:
            out.append(Run(node["raw"], bold, italic))
    return out


class _HtmlTables(HTMLParser):
    """문서 속 <table> 을 Block 으로 바꾼다. rowspan/colspan 을 풀어 모든 행의 칸 수를 맞추고, 병합 위치는 spans 에 남긴다."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.blocks: list[Block] = []
        self._rows: list[list[dict]] | None = None
        self._cell: dict | None = None

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag == "table":
            self._rows = []
        elif tag == "tr" and self._rows is not None:
            self._rows.append([])
        elif tag in ("td", "th") and self._rows:
            fill = ""
            if "bar" in (a.get("class") or "").split() or a.get("bgcolor"):
                fill = (a.get("bgcolor") or "#7055C8").lstrip("#").upper()
            self._cell = {
                "runs": [], "header": tag == "th", "fill": fill,
                "rowspan": max(1, int(a.get("rowspan") or 1)), "colspan": max(1, int(a.get("colspan") or 1)),
            }
            self._rows[-1].append(self._cell)
        elif tag == "br" and self._cell is not None:
            self._cell["runs"].append(Run("\n", self._cell["header"]))

    def handle_endtag(self, tag):
        if tag in ("td", "th"):
            self._cell = None
        elif tag == "table" and self._rows is not None:
            self.blocks.append(self._build(self._rows))
            self._rows = None

    def handle_data(self, data):
        if self._cell is not None and data.strip("\n\r\t") != "":
            text = " ".join(data.split()) if data.strip() else " "
            self._cell["runs"].append(Run(text.replace("\xa0", " ") if text.strip() else "", self._cell["header"]))

    @staticmethod
    def _build(rows: list[list[dict]]) -> Block:
        occupied: set[tuple[int, int]] = set()
        placed: dict[tuple[int, int], dict] = {}
        for r, row in enumerate(rows):
            c = 0
            for cell in row:
                while (r, c) in occupied:
                    c += 1
                placed[(r, c)] = cell
                for dr in range(cell["rowspan"]):
                    for dc in range(cell["colspan"]):
                        if dr or dc:
                            occupied.add((r + dr, c + dc))
                c += cell["colspan"]
        n_rows = len(rows)
        n_cols = max([c + cell["colspan"] for (r, c), cell in placed.items()] or [0])
        block = Block("table", rows=[[[] for _ in range(n_cols)] for _ in range(n_rows)], headers=set())
        for (r, c), cell in placed.items():
            if r >= n_rows:
                continue
            block.rows[r][c] = [run for run in cell["runs"] if run.text != ""]
            if cell["rowspan"] > 1 or cell["colspan"] > 1:
                block.spans[(r, c)] = (min(cell["rowspan"], n_rows - r), min(cell["colspan"], n_cols - c))
            if cell["header"]:
                block.headers.add((r, c))
            if cell["fill"]:
                block.fills[(r, c)] = cell["fill"]
        return block


def html_table_blocks(raw: str) -> list[Block]:
    parser = _HtmlTables()
    parser.feed(raw or "")
    parser.close()
    return parser.blocks


def markdown_blocks(markdown: str) -> list[Block]:
    blocks: list[Block] = []

    def walk_list(node: dict, depth: int) -> None:
        ordered = bool((node.get("attrs") or {}).get("ordered"))
        start = int((node.get("attrs") or {}).get("start") or 1)
        for index, item in enumerate(node.get("children", [])):
            for child in item.get("children", []):
                if child.get("type") in ("block_text", "paragraph"):
                    blocks.append(Block("list_item", _runs(child.get("children", [])), level=depth, ordered=ordered, number=start + index))
                elif child.get("type") == "list":
                    walk_list(child, depth + 1)

    for node in _MD(markdown or ""):
        kind = node.get("type")
        if kind == "heading":
            blocks.append(Block("heading", _runs(node.get("children", [])), level=int(node["attrs"]["level"])))
        elif kind == "paragraph":
            blocks.append(Block("paragraph", _runs(node.get("children", []))))
        elif kind == "list":
            walk_list(node, 0)
        elif kind == "table":
            rows: list[list[list[Run]]] = []
            for part in node.get("children", []):
                if part.get("type") == "table_head":
                    rows.append([_runs(cell.get("children", [])) for cell in part.get("children", [])])
                elif part.get("type") == "table_body":
                    for row in part.get("children", []):
                        rows.append([_runs(cell.get("children", [])) for cell in row.get("children", [])])
            blocks.append(Block("table", rows=rows))
        elif kind == "block_html":
            blocks.extend(html_table_blocks(str(node.get("raw") or "")))
        elif kind == "thematic_break":
            blocks.append(Block("rule"))
        elif kind == "block_quote":
            for child in node.get("children", []):
                blocks.append(Block("paragraph", _runs(child.get("children", []), italic=True)))
    return blocks


def _plain(runs: list[Run]) -> str:
    return "".join(r.text for r in runs)


# 일정표(월별 막대)의 막대 두께(pt). 칸 높이를 다 채우면 뚱뚱해 보여서 가운데에 얇게 그린다.
BAR_THICKNESS_PT = 4.5


def _display_width(text: str) -> int:
    return sum(2 if "\u1100" <= ch <= "\uffdc" else 1 for ch in text)


# --------------------------------------------------------------------------- Word
KR_FONT = "맑은 고딕"


def to_docx(markdown: str, *, title: str = "업무 인수인계서") -> bytes:
    from docx import Document
    from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_LINE_SPACING
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    doc.core_properties.title = title
    doc.core_properties.author = "AMIGO 인수인계 시스템"
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2.0)
        section.top_margin = section.bottom_margin = Cm(2.0)

    def set_font(style_or_run, size: float | None = None) -> None:
        font = style_or_run.font
        font.name = KR_FONT
        if size:
            font.size = Pt(size)
        rpr = style_or_run.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(attr), KR_FONT)
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rfonts.get(qn(attr)) is not None:
                del rfonts.attrib[qn(attr)]

    set_font(doc.styles["Normal"], 10)
    for name, size in (("Title", 22), ("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)):
        set_font(doc.styles[name], size)
        doc.styles[name].font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)

    def add_runs(paragraph, runs: list[Run], size: float | None = None) -> None:
        for run in runs:
            for i, piece in enumerate(run.text.split("\n")):
                if i:
                    paragraph.add_run().add_break()
                r = paragraph.add_run(piece)
                r.bold = run.bold or None
                r.italic = run.italic or None
                set_font(r, size)

    def shade(cell, color: str) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), color)
        tc_pr.append(shd)

    def thin_bar(cell, color: str) -> None:
        """칸을 칠하는 대신 칸 가운데에 얇은 막대(위쪽 문단 테두리)를 그린다. 이웃 칸과 이어지도록 좌우 여백은 0."""
        tc_pr = cell._tc.get_or_add_tcPr()
        margins = OxmlElement("w:tcMar")
        for side in ("left", "right"):
            edge = OxmlElement(f"w:{side}")
            edge.set(qn("w:w"), "0")
            edge.set(qn("w:type"), "dxa")
            margins.append(edge)
        tc_pr.insert_element_before(margins, "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark", "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge", "w:tcPrChange")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        paragraph = cell.paragraphs[0]
        fmt = paragraph.paragraph_format
        fmt.space_before = fmt.space_after = Pt(0)
        fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        fmt.line_spacing = Pt(1)
        border = OxmlElement("w:pBdr")
        top = OxmlElement("w:top")
        top.set(qn("w:val"), "single")
        top.set(qn("w:sz"), str(round(BAR_THICKNESS_PT * 8)))  # 1/8 pt 단위
        top.set(qn("w:space"), "0")
        top.set(qn("w:color"), color)
        border.append(top)
        paragraph._p.get_or_add_pPr().insert_element_before(
            border, "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE",
            "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
            "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle",
            "w:rPr", "w:sectPr", "w:pPrChange",
        )

    for block in markdown_blocks(markdown):
        if block.kind == "heading":
            style = "Title" if block.level == 1 else f"Heading {min(block.level - 1, 3)}"
            add_runs(doc.add_paragraph(style=style), block.runs)
        elif block.kind == "paragraph":
            add_runs(doc.add_paragraph(), block.runs)
        elif block.kind == "list_item" and block.ordered:
            # 근거 번호([1], [2] …)와 어긋나지 않도록 Word 자동 번호 대신 Markdown 의 번호를 그대로 쓴다
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.6 + 0.6 * block.level)
            paragraph.paragraph_format.first_line_indent = Cm(-0.6)
            add_runs(paragraph, [Run(f"{block.number}. ")] + block.runs)
        elif block.kind == "list_item":
            style = "List Bullet" if block.level == 0 else f"List Bullet {min(block.level + 1, 3)}"
            add_runs(doc.add_paragraph(style=style), block.runs)
        elif block.kind == "table" and block.rows:
            cols = max(len(r) for r in block.rows)
            table = doc.add_table(rows=len(block.rows), cols=cols)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            covered = {
                (r + dr, c + dc)
                for (r, c), (rs, cs) in block.spans.items()
                for dr in range(rs)
                for dc in range(cs)
                if dr or dc
            }
            for r_index, row in enumerate(block.rows):
                for c_index in range(cols):
                    if (r_index, c_index) in covered:
                        continue
                    cell = table.cell(r_index, c_index)
                    runs = row[c_index] if c_index < len(row) else []
                    if block.is_header(r_index, c_index):
                        runs = [Run(x.text, True, x.italic) for x in runs]
                        shade(cell, "E8EEF7")
                    if (r_index, c_index) in block.fills and not _plain(runs).strip():
                        thin_bar(cell, block.fills[(r_index, c_index)])  # 일정 막대(글자 없는 색칠 칸)
                        continue
                    if (r_index, c_index) in block.fills:
                        shade(cell, block.fills[(r_index, c_index)])
                    add_runs(cell.paragraphs[0], runs, size=9)
            for (r_index, c_index), (rs, cs) in block.spans.items():  # 병합은 내용을 채운 뒤에 한다
                table.cell(r_index, c_index).merge(table.cell(r_index + rs - 1, c_index + cs - 1))
            doc.add_paragraph()
        elif block.kind == "rule":
            doc.add_paragraph("―" * 30)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------- PDF
_FONT_CANDIDATES = [
    # (일반, 굵게, TTC 하위 글꼴 번호)
    ("/usr/share/fonts/truetype/nanum/NanumGothic.ttf", "/usr/share/fonts/truetype/nanum/NanumGothicBold.ttf", 0),
    ("C:/Windows/Fonts/malgun.ttf", "C:/Windows/Fonts/malgunbd.ttf", 0),
    ("/Library/Fonts/NanumGothic.ttf", "/Library/Fonts/NanumGothicBold.ttf", 0),
    ("/System/Library/Fonts/Supplemental/AppleGothic.ttf", None, 0),
    ("/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc", None, 0),
    ("/usr/share/fonts/truetype/unfonts-core/UnDotum.ttf", "/usr/share/fonts/truetype/unfonts-core/UnDotumBold.ttf", 0),
]
_REGISTERED: tuple[str, str] | None = None


def _pdf_fonts() -> tuple[str, str]:
    """(본문 글꼴, 굵은 글꼴) 이름을 등록하고 돌려준다."""
    global _REGISTERED
    if _REGISTERED:
        return _REGISTERED
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = list(_FONT_CANDIDATES)
    custom = os.getenv("AMIGO_PDF_FONT")
    if custom:
        candidates.insert(0, (custom, os.getenv("AMIGO_PDF_FONT_BOLD") or None, 0))
    for regular, bold, index in candidates:
        if not regular or not os.path.isfile(regular):
            continue
        try:
            pdfmetrics.registerFont(TTFont("AmigoKR", regular, subfontIndex=index))
            bold_name = "AmigoKR"
            if bold and os.path.isfile(bold):
                pdfmetrics.registerFont(TTFont("AmigoKR-Bold", bold, subfontIndex=index))
                bold_name = "AmigoKR-Bold"
            pdfmetrics.registerFontFamily("AmigoKR", normal="AmigoKR", bold=bold_name, italic="AmigoKR", boldItalic=bold_name)
            _REGISTERED = ("AmigoKR", bold_name)
            return _REGISTERED
        except Exception:  # CFF 기반 OTF 등 reportlab 이 못 읽는 글꼴은 건너뛴다
            continue
    pdfmetrics.registerFont(UnicodeCIDFont("HYGothic-Medium"))
    _REGISTERED = ("HYGothic-Medium", "HYGothic-Medium")
    return _REGISTERED


def _markup(runs: list[Run]) -> str:
    parts = []
    for run in runs:
        text = escape(run.text).replace("\n", "<br/>")
        if run.bold:
            text = f"<b>{text}</b>"
        parts.append(text)
    return "".join(parts)


def to_pdf(markdown: str, *, title: str = "업무 인수인계서") -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Flowable, HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    class ThinBar(Flowable):
        """칸 폭만큼 이어지는 얇은 막대(칸 가운데에 놓인다)."""

        def __init__(self, color) -> None:
            super().__init__()
            self.color, self.width = color, 0

        def wrap(self, avail_width, avail_height):
            self.width = avail_width
            return avail_width, BAR_THICKNESS_PT

        def draw(self) -> None:
            self.canv.setFillColor(self.color)
            self.canv.rect(0, 0, self.width, BAR_THICKNESS_PT, stroke=0, fill=1)

    regular, bold = _pdf_fonts()
    navy = colors.HexColor("#1F2A44")
    body = ParagraphStyle("body", fontName=regular, fontSize=9.5, leading=14.5, wordWrap="CJK", alignment=TA_LEFT, spaceAfter=4)
    cell = ParagraphStyle("cell", parent=body, fontSize=8.5, leading=12, spaceAfter=0)
    head_cell = ParagraphStyle("head_cell", parent=cell, fontName=bold)
    headings = {
        1: ParagraphStyle("h1", parent=body, fontName=bold, fontSize=18, leading=24, textColor=navy, spaceAfter=10),
        2: ParagraphStyle("h2", parent=body, fontName=bold, fontSize=13, leading=18, textColor=navy, spaceBefore=10, spaceAfter=6),
        3: ParagraphStyle("h3", parent=body, fontName=bold, fontSize=11, leading=16, textColor=navy, spaceBefore=6, spaceAfter=4),
    }

    page_width = A4[0] - 36 * mm
    story = []
    for block in markdown_blocks(markdown):
        if block.kind == "heading":
            story.append(Paragraph(_markup(block.runs), headings[min(block.level, 3)]))
            if block.level == 2:
                story.append(HRFlowable(width="100%", thickness=0.6, color=colors.HexColor("#C9D3E3"), spaceAfter=4))
        elif block.kind == "paragraph":
            story.append(Paragraph(_markup(block.runs), body))
        elif block.kind == "list_item":
            bullet = f"{block.number}." if block.ordered else "•"
            style = ParagraphStyle("li", parent=body, leftIndent=12 + 12 * block.level, bulletIndent=2 + 12 * block.level)
            story.append(Paragraph(_markup(block.runs), style, bulletText=bullet))
        elif block.kind == "table" and block.rows:
            cols = max(len(r) for r in block.rows)
            merged_across = {(r, c) for (r, c), (rs, cs) in block.spans.items() if cs > 1}
            # 열 너비: 한글은 영문보다 두 배 넓으므로 표시 폭 기준으로 나누고, 짧은 열도 최소 4글자 폭은 보장한다(여러 열에 걸친 칸은 제외)
            lengths = [
                max((_display_width(_plain(r[c])) for ri, r in enumerate(block.rows) if c < len(r) and (ri, c) not in merged_across), default=0)
                for c in range(cols)
            ]
            weights = [min(max(n, 9), 70) for n in lengths]
            total = sum(weights) or 1
            widths = [page_width * w / total for w in weights]
            data = []
            bars = {key for key in block.fills if not _plain(block.rows[key[0]][key[1]] if key[1] < len(block.rows[key[0]]) else []).strip()}
            for r_index, row in enumerate(block.rows):
                data.append(
                    [
                        ThinBar(colors.HexColor("#" + block.fills[(r_index, c)]))
                        if (r_index, c) in bars
                        else Paragraph(_markup(row[c]) if c < len(row) else "", head_cell if block.is_header(r_index, c) else cell)
                        for c in range(cols)
                    ]
                )
            table = Table(data, colWidths=widths, repeatRows=1)
            commands = [
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#9AA8BD")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ]
            if block.headers is None:
                commands.append(("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF7")))
            else:
                commands += [("BACKGROUND", (c, r), (c, r), colors.HexColor("#E8EEF7")) for (r, c) in sorted(block.headers)]
            for (r, c), (rs, cs) in block.spans.items():
                commands.append(("SPAN", (c, r), (c + cs - 1, r + rs - 1)))
            for (r, c), color in block.fills.items():
                rs, cs = block.spans.get((r, c), (1, 1))
                if (r, c) in bars:  # 얇은 막대: 이웃 칸과 이어지게 좌우 여백을 없애고 세로 가운데에 둔다
                    commands += [
                        ("LEFTPADDING", (c, r), (c + cs - 1, r + rs - 1), 0),
                        ("RIGHTPADDING", (c, r), (c + cs - 1, r + rs - 1), 0),
                        ("VALIGN", (c, r), (c + cs - 1, r + rs - 1), "MIDDLE"),
                    ]
                else:
                    commands.append(("BACKGROUND", (c, r), (c + cs - 1, r + rs - 1), colors.HexColor("#" + color)))
            table.setStyle(TableStyle(commands))
            story += [table, Spacer(1, 6)]
        elif block.kind == "rule":
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.grey))

    def footer(canvas, doc_template) -> None:
        canvas.saveState()
        canvas.setFont(regular, 8)
        canvas.setFillColor(colors.grey)
        canvas.drawString(18 * mm, 10 * mm, f"{title} · AMIGO 자동 작성 · {date.today().isoformat()}")
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"{doc_template.page}")
        canvas.restoreState()

    buffer = io.BytesIO()
    SimpleDocTemplate(
        buffer, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=18 * mm, title=title, author="AMIGO"
    ).build(story or [Paragraph("(내용 없음)", body)], onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()
